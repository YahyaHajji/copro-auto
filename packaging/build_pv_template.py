from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
RUNTIME = ROOT / "resources" / "templates" / "runtime"


def _merge_styles(cover, narrative) -> None:
    target = narrative.styles.element
    existing = {style.get(qn("w:styleId")) for style in target.findall(qn("w:style"))}
    for style in cover.styles.element.findall(qn("w:style")):
        style_id = style.get(qn("w:styleId"))
        if style_id not in existing:
            target.append(deepcopy(style))
            existing.add(style_id)


def _merge_numbering(cover, narrative) -> dict[str, str]:
    source = cover.part.numbering_part.element
    target = narrative.part.numbering_part.element
    target_abstract_ids = [
        int(item.get(qn("w:abstractNumId"))) for item in target.findall(qn("w:abstractNum"))
    ]
    target_num_ids = [int(item.get(qn("w:numId"))) for item in target.findall(qn("w:num"))]
    abstract_offset = max(target_abstract_ids, default=-1) + 1
    num_offset = max(target_num_ids, default=0) + 1
    abstract_map: dict[str, str] = {}
    num_map: dict[str, str] = {}

    for abstract in source.findall(qn("w:abstractNum")):
        old_id = abstract.get(qn("w:abstractNumId"))
        new_id = str(abstract_offset + len(abstract_map))
        copied = deepcopy(abstract)
        copied.set(qn("w:abstractNumId"), new_id)
        target.insert(len(target.findall(qn("w:abstractNum"))), copied)
        abstract_map[old_id] = new_id

    for numbering in source.findall(qn("w:num")):
        old_id = numbering.get(qn("w:numId"))
        new_id = str(num_offset + len(num_map))
        copied = deepcopy(numbering)
        copied.set(qn("w:numId"), new_id)
        abstract_id = copied.find(qn("w:abstractNumId"))
        if abstract_id is not None:
            old_abstract_id = abstract_id.get(qn("w:val"))
            abstract_id.set(qn("w:val"), abstract_map[old_abstract_id])
        target.append(copied)
        num_map[old_id] = new_id
    return num_map


def _copy_cover_body(cover, narrative, num_map: dict[str, str]) -> None:
    first_narrative_element = narrative.element.body[0]
    for element in cover.element.body:
        if element.tag == qn("w:sectPr"):
            continue
        copied = deepcopy(element)
        for num_id in copied.iter(qn("w:numId")):
            old_id = num_id.get(qn("w:val"))
            if old_id in num_map:
                num_id.set(qn("w:val"), num_map[old_id])
        first_narrative_element.addprevious(copied)


def _remove_trailing_empty_section(document) -> None:
    """Drop the unused final PV2 section that otherwise renders as a blank page."""
    paragraphs = document.paragraphs
    section_break_index = next(
        (
            index
            for index in range(len(paragraphs) - 1, -1, -1)
            if paragraphs[index]._p.pPr is not None
            and paragraphs[index]._p.pPr.sectPr is not None
        ),
        None,
    )
    if section_break_index is None:
        return
    trailing = paragraphs[section_break_index + 1:]
    if any(paragraph.text.strip() for paragraph in trailing):
        return

    break_paragraph = paragraphs[section_break_index]
    section_properties = deepcopy(break_paragraph._p.pPr.sectPr)
    body = document.element.body
    for paragraph in trailing:
        body.remove(paragraph._p)
    body.remove(break_paragraph._p)
    final_section = body.find(qn("w:sectPr"))
    if final_section is not None:
        body.remove(final_section)
    body.append(section_properties)


def build() -> Path:
    cover = Document(RUNTIME / "pv_division_1.docx")
    narrative = Document(RUNTIME / "pv_division_2.docx")
    _remove_trailing_empty_section(narrative)
    _merge_styles(cover, narrative)
    num_map = _merge_numbering(cover, narrative)
    _copy_cover_body(cover, narrative, num_map)
    first_narrative_element = next(
        paragraph._p for paragraph in narrative.paragraphs
        if paragraph.text.strip().startswith("Après avoir reconnu")
    )

    break_paragraph = OxmlElement("w:p")
    paragraph_properties = OxmlElement("w:pPr")
    section_properties = deepcopy(cover.sections[-1]._sectPr)
    section_type = section_properties.find(qn("w:type"))
    if section_type is None:
        section_type = OxmlElement("w:type")
        section_properties.insert(0, section_type)
    section_type.set(qn("w:val"), "nextPage")
    paragraph_properties.append(section_properties)
    break_paragraph.append(paragraph_properties)
    first_narrative_element.addprevious(break_paragraph)

    destination = RUNTIME / "pv_division.docx"
    narrative.save(destination)
    return destination


if __name__ == "__main__":
    print(build())
