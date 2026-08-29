from __future__ import annotations

import hashlib
import logging
import sys
from collections import defaultdict
from copy import deepcopy
from decimal import Decimal
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from copro_auto.domain.calculations import calculate_shares
from copro_auto.domain.models import Level, PartNature, Project

from .mappings import (
    base_replacements,
    client_sentence,
    document_replacements,
    format_decimal,
    format_level_heading,
    level_description,
    part_description,
)


LOGGER = logging.getLogger(__name__)


def resource_root() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys._MEIPASS) / "resources"  # type: ignore[attr-defined]
    return Path(__file__).resolve().parents[3] / "resources"


def _enable_field_updates(document: DocumentObject) -> None:
    """Ask Word to refresh page-count and other fields when the file opens."""
    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def _trim_trailing_empty_paragraphs(document: DocumentObject) -> None:
    """Remove template-only empty paragraphs that spill onto a blank final page."""
    body = document.element.body
    for paragraph in reversed(document.paragraphs):
        if paragraph.text.strip():
            break
        properties = paragraph._p.pPr
        if properties is not None and properties.sectPr is not None:
            break
        if paragraph._p.xpath(".//w:drawing | .//w:pict | .//w:br"):
            break
        body.remove(paragraph._p)


def _ensure_footer_field_clearance(document: DocumentObject) -> None:
    """Keep refreshed two-digit page counts inside mirrored page margins."""
    for section in document.sections:
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            for paragraph in footer.paragraphs:
                if paragraph._p.xpath(".//w:fldSimple | .//w:instrText"):
                    clearance = 0.75 if paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT else 0.35
                    paragraph.paragraph_format.right_indent = Inches(clearance)


def _set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _replace_in_paragraph(paragraph, replacements: dict[str, str]) -> None:
    for source, target in replacements.items():
        if not source or source == target:
            continue
        search_from = 0
        while True:
            full_text = paragraph.text
            start = full_text.find(source, search_from)
            if start < 0:
                break
            end = start + len(source)
            cursor = 0
            start_run = start_offset = end_run = end_offset = None
            for run_index, run in enumerate(paragraph.runs):
                run_end = cursor + len(run.text)
                if start_run is None and start < run_end:
                    start_run = run_index
                    start_offset = start - cursor
                if end_run is None and end <= run_end:
                    end_run = run_index
                    end_offset = end - cursor
                    break
                cursor = run_end
            if start_run is None or end_run is None or start_offset is None or end_offset is None:
                _set_paragraph_text(paragraph, full_text.replace(source, target, 1))
                search_from = start + len(target)
                continue
            first = paragraph.runs[start_run]
            if start_run == end_run:
                first.text = first.text[:start_offset] + target + first.text[end_offset:]
                search_from = start + len(target)
                continue
            last = paragraph.runs[end_run]
            first.text = first.text[:start_offset] + target
            for run in paragraph.runs[start_run + 1:end_run]:
                run.text = ""
            last.text = last.text[end_offset:]
            search_from = start + len(target)


def replace_everywhere(document: DocumentObject, replacements: dict[str, str]) -> None:
    seen = set()

    def replace_story(root, parent) -> None:
        for paragraph_element in root.iter(qn("w:p")):
            if paragraph_element in seen:
                continue
            seen.add(paragraph_element)
            _replace_in_paragraph(Paragraph(paragraph_element, parent), replacements)

    replace_story(document.element.body, document)
    for section in document.sections:
        for part in (section.header, section.footer):
            replace_story(part._element, part)


def _clear_after(table: Table, keep_rows: int) -> None:
    for row in list(table.rows)[keep_rows:]:
        table._tbl.remove(row._tr)


def _set_cell_text(cell: _Cell, value: str) -> None:
    for index, paragraph in enumerate(cell.paragraphs):
        _set_paragraph_text(paragraph, value if index == 0 else "")


def _add_values(table: Table, values: list[str], prototype=None):
    if prototype is None:
        row = table.add_row()
    else:
        table._tbl.append(deepcopy(prototype))
        row = table.rows[-1]
    cells = list(row.cells)
    seen: set[object] = set()
    for cell in cells:
        identity = cell._tc
        if identity not in seen:
            _set_cell_text(cell, "")
            seen.add(identity)
    assigned: set[object] = set()
    for index, value in enumerate(values[: len(cells)]):
        identity = cells[index]._tc
        if identity not in assigned:
            _set_cell_text(cells[index], value)
            assigned.add(identity)
    return row


def _populate_table_a(document: DocumentObject, project: Project) -> None:
    table = document.tables[0]
    data_prototype = deepcopy(table.rows[1]._tr)
    _clear_after(table, 1)
    private_number = 1
    for level in sorted(project.levels, key=lambda item: item.order):
        for part in level.parts:
            if part.nature is not PartNature.PRIVATE:
                continue
            _add_values(table, [
                f"{project.identity.property_name}-{private_number}", project.identity.land_title, part.index,
                format_decimal(part.surfaces.cadastral_total), level.name, part.consistency, part.observations,
            ], data_prototype)
            private_number += 1


def _populate_table_b(document: DocumentObject, project: Project) -> None:
    table = document.tables[0]
    heading_prototype = deepcopy(table.rows[2]._tr)
    private_prototype = deepcopy(table.rows[3]._tr)
    common_prototype = deepcopy(table.rows[4]._tr)
    total_prototype = deepcopy(table.rows[6]._tr)
    _clear_after(table, 2)
    private_number = 1
    for level in sorted(project.levels, key=lambda item: item.order):
        _add_values(table, [format_level_heading(level)], heading_prototype)
        for part in level.parts:
            private_index = part.index if part.nature is PartNature.PRIVATE else ""
            common_index = part.index if part.nature is PartNature.COMMON else ""
            if part.nature is PartNature.PRIVATE:
                name = f"{project.identity.property_name}-{private_number}"
                private_number += 1
            else:
                name = project.identity.property_name
            _add_values(table, [
                name, f"T.{project.identity.land_title}" if part.nature is PartNature.COMMON else "",
                private_index, common_index, format_decimal(part.surfaces.inside_title),
                format_decimal(part.surfaces.cadastral_total), part.consistency, part.observations,
            ], private_prototype if part.nature is PartNature.PRIVATE else common_prototype)
        totals_inside = sum((part.surfaces.inside_title for part in level.parts), Decimal("0"))
        totals_all = sum((part.surfaces.cadastral_total for part in level.parts), Decimal("0"))
        _add_values(table, ["Total", "", "", "", format_decimal(totals_inside), format_decimal(totals_all), "", ""], total_prototype)


def _component_rows(part) -> list[tuple[str, tuple[Decimal, Decimal, Decimal, Decimal, Decimal]]]:
    surfaces = part.surfaces
    excluding_balcony = surfaces.excluding_balcony or surfaces.cadastral_total
    main_base = excluding_balcony - surfaces.courtyard - surfaces.garage - surfaces.terrace
    if main_base < 0:
        main_base = excluding_balcony
    main_total = main_base + surfaces.courtyard + surfaces.balcony + surfaces.terrace
    consistency = part.consistency.replace("+ Garage", "").replace("+Garage", "").strip()
    if consistency.casefold() == "appartement":
        consistency = "Appartement (Habitation)"
    rows = [(consistency or part.consistency, (
        main_base, surfaces.courtyard, surfaces.balcony, surfaces.terrace, main_total,
    ))]
    if surfaces.garage:
        rows.append(("Garage", (
            surfaces.garage, Decimal("0"), Decimal("0"), Decimal("0"), surfaces.garage,
        )))
    return rows


def _populate_recap(document: DocumentObject, project: Project) -> None:
    detailed = document.tables[1]
    detailed_prototype = deepcopy(detailed.rows[2]._tr)
    _clear_after(detailed, 2)
    aggregate: dict[str, list[Decimal]] = defaultdict(lambda: [Decimal("0") for _ in range(5)])
    for level in sorted(project.levels, key=lambda item: item.order):
        for part in level.parts:
            if part.nature is not PartNature.PRIVATE:
                continue
            for consistency, values in _component_rows(part):
                _add_values(detailed, [level.name, consistency, *[format_decimal(value) if value else "" for value in values]], detailed_prototype)
                for index, value in enumerate(values):
                    aggregate[consistency][index] += value
    totals = document.tables[3]
    total_prototype = deepcopy(totals.rows[2]._tr)
    _clear_after(totals, 2)
    for consistency, values in sorted(aggregate.items()):
        label = consistency if consistency == "Garage" else f"Total/{consistency}"
        _add_values(totals, [label, *[format_decimal(value) if value else "" for value in values]], total_prototype)


def _populate_regulation_tables(document: DocumentObject, project: Project) -> None:
    shares = {share.part_id: share for share in calculate_shares(project)}
    distribution = _find_table(document, "Indice des parties")
    heading_prototype = deepcopy(distribution.rows[3]._tr)
    private_prototype = deepcopy(distribution.rows[4]._tr)
    common_prototype = deepcopy(distribution.rows[5]._tr)
    total_prototype = deepcopy(distribution.rows[7]._tr)
    general_prototype = deepcopy(distribution.rows[-1]._tr)
    _clear_after(distribution, 3)
    for level in sorted(project.levels, key=lambda item: item.order):
        _add_values(distribution, [format_level_heading(level)], heading_prototype)
        for part in level.parts:
            share = shares.get(part.id)
            _add_values(distribution, [
                part.index if part.nature is PartNature.PRIVATE else "",
                part.index if part.nature is PartNature.COMMON else "",
                part.consistency, format_decimal(part.surfaces.inside_title),
                format_decimal(part.surfaces.cadastral_total),
                "" if share is None else format_decimal(share.percentage),
                "" if share is None else str(share.ten_thousandths), part.observations,
            ], private_prototype if part.nature is PartNature.PRIVATE else common_prototype)
        totals_inside = sum((part.surfaces.inside_title for part in level.parts), Decimal("0"))
        totals_all = sum((part.surfaces.cadastral_total for part in level.parts), Decimal("0"))
        level_shares = [shares[part.id] for part in level.parts if part.id in shares]
        _add_values(distribution, [
            "Total", "Total", "Total", format_decimal(totals_inside), format_decimal(totals_all),
            format_decimal(sum((share.percentage for share in level_shares), Decimal("0"))) if level_shares else "",
            str(sum(share.ten_thousandths for share in level_shares)) if level_shares else "", "",
        ], total_prototype)
    private = [part for level in project.levels for part in level.parts if part.nature is PartNature.PRIVATE]
    _add_values(distribution, [
        "Totaux généraux", "Totaux généraux", "Totaux généraux",
        format_decimal(sum((part.surfaces.inside_title for part in private), Decimal("0"))),
        format_decimal(sum((part.surfaces.cadastral_total for part in private), Decimal("0"))),
        "100", "10000", "",
    ], general_prototype)
    voices = _find_table(document, "N° d’ordre")
    voice_prototype = deepcopy(voices.rows[1]._tr)
    voice_total_prototype = deepcopy(voices.rows[-1]._tr)
    _clear_after(voices, 1)
    order = 1
    for level in sorted(project.levels, key=lambda item: item.order):
        for part in level.parts:
            share = shares.get(part.id)
            if share is None:
                continue
            _add_values(voices, [str(order), part.index, level.name, part.consistency, format_decimal(share.reference_surface), format_decimal(share.percentage)], voice_prototype)
            order += 1
    _add_values(voices, ["", "", "", "Total", format_decimal(sum((s.reference_surface for s in shares.values()), Decimal("0"))), "100.00"], voice_total_prototype)


def _populate_pv_overview(document: DocumentObject, project: Project) -> None:
    identity = project.identity
    private_indices = {
        part.index.casefold() for level in project.levels for part in level.parts
        if part.nature is PartNature.PRIVATE
    }
    common_indices = {
        part.index.casefold() for level in project.levels for part in level.parts
        if part.nature is PartNature.COMMON
    }
    for paragraph in document.paragraphs:
        normalized = paragraph.text.strip().casefold()
        if normalized.startswith("après avoir reconnu"):
            _set_paragraph_text(paragraph, (
                "Après avoir reconnu la consistance générale de l'immeuble à diviser qui comprend une construction "
                f"composée de {identity.overall_consistency}, d’une hauteur totale de {format_decimal(identity.total_height)} m "
                "au-dessus du niveau du trottoir, l'immeuble est édifié sur la totalité de la propriété dite « "
                f"{identity.property_name} », objet du titre foncier N° {identity.land_title} d’une surface de "
                f"{format_decimal(identity.land_area)} m², ainsi qu’il est indiqué au plan de mise en concordance."
            ))
        elif normalized.startswith("avons procédé à la division"):
            total = len(private_indices) + len(common_indices)
            _set_paragraph_text(paragraph, (
                f"Avons procédé à la division du dit immeuble qui est divisé en {total} parties : "
                f"{len(private_indices)} parties privatives et {len(common_indices)} parties communes, de la façon suivante :"
            ))


def _find_table(document: DocumentObject, needle: str) -> Table:
    normalized = needle.casefold()
    for table in document.tables:
        if any(normalized in cell.text.casefold() for row in table.rows for cell in row.cells):
            return table
    raise ValueError(f"Table de modèle introuvable : {needle}")


def _set_first_matching_paragraph(document: DocumentObject, prefix: str, text: str) -> bool:
    normalized = prefix.casefold()
    for paragraph in document.paragraphs:
        if paragraph.text.strip().casefold().startswith(normalized):
            _set_paragraph_text(paragraph, text)
            return True
    return False


def _populate_client_paragraph(document: DocumentObject, project: Project) -> None:
    if not _set_first_matching_paragraph(document, "MM. (4)", client_sentence(project)):
        LOGGER.warning("client_paragraph_not_found")


def _populate_regulation_overview(document: DocumentObject, project: Project) -> None:
    identity = project.identity
    _populate_client_paragraph(document, project)
    if len(document.tables) >= 2:
        identity_table = document.tables[0]
        if len(identity_table.rows) > 3:
            _set_cell_text(identity_table.rows[3].cells[2], identity.overall_consistency)
        boundaries = document.tables[1]
        keys = ("northeast", "northwest", "southeast", "southwest")
        for row, key in zip(boundaries.rows, keys, strict=False):
            _set_cell_text(row.cells[2], identity.boundaries.get(key, ""))
    _set_first_matching_paragraph(
        document,
        "Les constructions sont édifiées",
        (
            "Les constructions sont édifiées sur la totalité de la propriété se composant d’un "
            f"immeuble ayant une consistance de {identity.overall_consistency}, d’une hauteur totale "
            f"de {format_decimal(identity.total_height)} m au-dessus du niveau du trottoir."
        ),
    )
    private_indices = {
        part.index.strip().casefold() for level in project.levels for part in level.parts
        if part.nature is PartNature.PRIVATE
    }
    common_indices = {
        part.index.strip().casefold() for level in project.levels for part in level.parts
        if part.nature is PartNature.COMMON
    }
    private_count = len(private_indices)
    common_count = len(common_indices)
    _set_first_matching_paragraph(
        document,
        "L’immeuble est divisé",
        (
            f"L’immeuble est divisé en {private_count + common_count} parties : "
            f"{private_count} parties privatives et {common_count} parties communes détaillées comme suit :"
        ),
    )


def _remove_body_range(start, end) -> object:
    parent = start.getparent()
    anchor = end.getnext()
    node = start
    while node is not None:
        next_node = node.getnext()
        parent.remove(node)
        if node is end:
            break
        node = next_node
    if anchor is None:
        raise ValueError("Ancre de narration introuvable après le dernier bloc de niveau.")
    return anchor


def _insert_level_blocks(
    document: DocumentObject,
    project: Project,
    *,
    start_table_index: int,
    end_table_index: int,
    heading_table_index: int,
    parts_table_index: int,
    description_prefix: str,
) -> None:
    tables = list(document.tables)
    heading_prototype = deepcopy(tables[heading_table_index]._tbl)
    parts_prototype = deepcopy(tables[parts_table_index]._tbl)
    description = next(
        paragraph for paragraph in document.paragraphs
        if paragraph.text.strip().casefold().startswith(description_prefix.casefold())
    )
    description_prototype = deepcopy(description._p)
    anchor = _remove_body_range(tables[start_table_index]._tbl, tables[end_table_index]._tbl)
    for level in sorted(project.levels, key=lambda item: item.order):
        heading_element = deepcopy(heading_prototype)
        heading_table = Table(heading_element, document)
        _clear_after(heading_table, 1)
        _set_cell_text(heading_table.rows[0].cells[0], level.name)
        anchor.addprevious(heading_element)

        description_element = deepcopy(description_prototype)
        description_paragraph = Paragraph(description_element, document)
        _set_paragraph_text(description_paragraph, level_description(level))
        anchor.addprevious(description_element)

        parts_element = deepcopy(parts_prototype)
        parts_table = Table(parts_element, document)
        _clear_after(parts_table, 1)
        _replace_cell_parts(parts_table.rows[0].cells[0], level.parts)
        anchor.addprevious(parts_element)

        spacer_element = deepcopy(description_prototype)
        _set_paragraph_text(Paragraph(spacer_element, document), "")
        anchor.addprevious(spacer_element)
def _write_part_paragraph(paragraph, part) -> None:
    prototype_runs = list(paragraph.runs)
    label_properties = deepcopy(prototype_runs[0]._r.rPr) if prototype_runs and prototype_runs[0]._r.rPr is not None else None
    detail_source = next((run for run in prototype_runs if not run.bold and run.text.strip()), None)
    detail_properties = (
        deepcopy(detail_source._r.rPr)
        if detail_source is not None and detail_source._r.rPr is not None
        else None
    )
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    text = part_description(part)
    label, separator, detail = text.partition(" : ")
    label_run = paragraph.add_run(label + (" :" if separator else ""))
    if label_properties is not None:
        label_run._r.insert(0, label_properties)
    if part.nature is PartNature.PRIVATE:
        label_run.bold = True
        label_run.underline = True
    if separator:
        detail_run = paragraph.add_run(" " + detail)
        if detail_properties is not None:
            detail_run._r.insert(0, detail_properties)


def _replace_cell_parts(cell: _Cell, parts) -> None:
    paragraph_properties = [
        deepcopy(paragraph._p.pPr) for paragraph in cell.paragraphs if paragraph._p.pPr is not None
    ]
    prototype_runs = next((paragraph.runs for paragraph in cell.paragraphs if paragraph.runs), [])
    label_run_properties = (
        deepcopy(prototype_runs[0]._r.rPr)
        if prototype_runs and prototype_runs[0]._r.rPr is not None else None
    )
    detail_source = next((run for run in prototype_runs if not run.bold and run.text.strip()), None)
    detail_run_properties = (
        deepcopy(detail_source._r.rPr)
        if detail_source is not None and detail_source._r.rPr is not None else None
    )
    for child in list(cell._tc):
        if child.tag != qn("w:tcPr"):
            cell._tc.remove(child)
    if not parts:
        cell.add_paragraph()
        return
    for index, part in enumerate(parts):
        paragraph = cell.add_paragraph()
        if paragraph_properties:
            paragraph._p.insert(0, deepcopy(paragraph_properties[min(index, len(paragraph_properties) - 1)]))
        label_seed = paragraph.add_run("PARTIE")
        if label_run_properties is not None:
            label_seed._r.insert(0, deepcopy(label_run_properties))
        detail_seed = paragraph.add_run(" détail")
        if detail_run_properties is not None:
            detail_seed._r.insert(0, deepcopy(detail_run_properties))
        _write_part_paragraph(paragraph, part)


def _populate_pv2_narrative(document: DocumentObject, project: Project) -> None:
    """Replace all historical fixed slots with one block per real level."""
    _insert_level_blocks(
        document,
        project,
        start_table_index=0,
        end_table_index=6,
        heading_table_index=2,
        parts_table_index=1,
        description_prefix="Le Premier",
    )


def _populate_regulation_narrative(document: DocumentObject, project: Project) -> None:
    """Replace all historical fixed slots with one block per real level."""
    _insert_level_blocks(
        document,
        project,
        start_table_index=2,
        end_table_index=8,
        heading_table_index=4,
        parts_table_index=3,
        description_prefix="Le Premier",
    )


def render_document(template: Path, destination: Path, project: Project, kind: str) -> Path:
    destination.parent.mkdir(parents=True, exist_ok=True)
    document = Document(template)
    replace_everywhere(document, base_replacements(project))
    replace_everywhere(document, document_replacements(project, kind))
    if kind == "tableau_a":
        _populate_table_a(document, project)
    elif kind == "tableau_b":
        _populate_table_b(document, project)
    elif kind == "tableau_recapitulatif":
        _populate_recap(document, project)
    elif kind == "reglement":
        _populate_regulation_overview(document, project)
        _populate_regulation_narrative(document, project)
        _populate_regulation_tables(document, project)
    elif kind == "pv_division":
        _populate_client_paragraph(document, project)
        _populate_pv_overview(document, project)
        _populate_pv2_narrative(document, project)
    if kind == "reglement":
        _trim_trailing_empty_paragraphs(document)
        _ensure_footer_field_clearance(document)
    _enable_field_updates(document)
    document.save(destination)
    LOGGER.info("document_rendered kind=%s file=%s", kind, destination.name)
    return destination


def template_hashes(template_dir: Path) -> dict[str, str]:
    return {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in sorted(template_dir.glob("*.docx"))}
