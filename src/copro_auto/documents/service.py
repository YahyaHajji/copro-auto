from __future__ import annotations

import os
import shutil
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document

from copro_auto import __version__
from copro_auto.domain.models import GenerationRecord, Project
from copro_auto.domain.validation import has_errors, validate_project

from .docx_renderer import render_document, resource_root, template_hashes


OUTPUTS = {
    "pv_division_1": "PV_Division_1.docx",
    "pv_division_2": "PV_Division_2.docx",
    "reglement": "Reglement_Copropriete.docx",
    "tableau_a": "Tableau_A.docx",
    "tableau_b": "Tableau_B.docx",
    "tableau_recapitulatif": "Tableau_Recapitulatif.docx",
}


class GenerationError(RuntimeError):
    pass


def _critical_text(path: Path) -> str:
    document = Document(path)
    chunks = [paragraph.text for paragraph in document.paragraphs]
    chunks.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    return "\n".join(chunks)


def _verify(path: Path, project: Project, kind: str) -> None:
    try:
        with ZipFile(path) as archive:
            if "word/document.xml" not in archive.namelist():
                raise GenerationError(f"Document Word incomplet : {path.name}")
    except (OSError, BadZipFile) as exc:
        raise GenerationError(f"Document Word invalide : {path.name}") from exc
    text = _critical_text(path)
    required = [project.identity.land_title]
    if kind != "tableau_recapitulatif":
        required.append(project.identity.property_name)
    for value in required:
        if value and value.casefold() not in text.casefold():
            raise GenerationError(f"{path.name} ne contient pas la valeur critique « {value} ».")


class DocumentGenerationService:
    def __init__(self, templates: Path | None = None) -> None:
        self.templates = templates or resource_root() / "templates" / "runtime"

    def generate(self, project: Project, output_dir: str | Path) -> list[Path]:
        issues = validate_project(project)
        if has_errors(issues):
            messages = "; ".join(issue.message for issue in issues if issue.severity.value == "error")
            raise GenerationError(f"Corrigez les erreurs avant génération : {messages}")
        output = Path(output_dir)
        output.mkdir(parents=True, exist_ok=True)
        temporary = Path(tempfile.mkdtemp(prefix="CoproAuto-generate-", dir=output.parent))
        generated: list[Path] = []
        try:
            for kind, filename in OUTPUTS.items():
                template = self.templates / f"{kind}.docx"
                if not template.exists():
                    raise GenerationError(f"Modèle introuvable : {template.name}")
                staged = render_document(template, temporary / filename, project, kind)
                _verify(staged, project, kind)
            for staged in temporary.glob("*.docx"):
                final = output / staged.name
                os.replace(staged, final)
                generated.append(final)
        except Exception:
            for final in generated:
                final.unlink(missing_ok=True)
            raise
        finally:
            shutil.rmtree(temporary, ignore_errors=True)
        project.generations.append(GenerationRecord(
            generated_at=datetime.now(timezone.utc).isoformat(),
            app_version=__version__,
            template_hashes=template_hashes(self.templates),
            files=[path.name for path in sorted(generated)],
            validation_ok=True,
        ))
        return sorted(generated)
