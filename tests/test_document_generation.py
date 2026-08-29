from copy import deepcopy
from decimal import Decimal
from zipfile import ZipFile

import pytest
from docx import Document
from lxml import etree

from copro_auto.documents.docx_renderer import resource_root
from copro_auto.documents.mappings import level_description
from copro_auto.documents.service import DocumentGenerationService, GenerationError, _verify


EXPECTED_OUTPUTS = {
    "PV_Division.docx",
    "Reglement_Copropriete.docx",
    "Tableau_A.docx",
    "Tableau_B.docx",
    "Tableau_Recapitulatif.docx",
}


def all_text(path):
    document = Document(path)
    return "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )


def package_text(path):
    chunks = []
    namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if not (name.startswith("word/") and name.endswith(".xml")):
                continue
            root = etree.fromstring(archive.read(name))
            for paragraph in root.xpath(".//w:p", namespaces=namespaces):
                text = "".join(paragraph.xpath(".//w:t/text()", namespaces=namespaces))
                if text:
                    chunks.append(text)
    return "\n".join(chunks)


def assert_french_characters_are_intact(path):
    text = package_text(path)
    for marker in ("?", "�", "Ã", "Â"):
        assert marker not in text, f"Caractère d’encodage invalide {marker!r} dans {path.name}"


def _heading_cells(document):
    return {
        cell.text.strip()
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        if len(row.cells) == 1
    }


def test_generates_five_documents_with_critical_values(tmp_path, yasmin_project):
    files = DocumentGenerationService().generate(yasmin_project, tmp_path)
    assert {path.name for path in files} == EXPECTED_OUTPUTS
    for path in files:
        text = package_text(path)
        assert_french_characters_are_intact(path)
        assert "119753/59" in text
        if path.name != "Tableau_Recapitulatif.docx":
            assert "yasmin 71" in text.casefold()
    regulation = all_text(tmp_path / "Reglement_Copropriete.docx")
    for value in ("3522", "3239", "10000"):
        assert value in regulation
    assert yasmin_project.generations[-1].validation_ok
    assert yasmin_project.generations[-1].files == sorted(EXPECTED_OUTPUTS)


def test_verification_rejects_question_mark_encoding_damage(tmp_path, yasmin_project):
    damaged = tmp_path / "damaged.docx"
    document = Document()
    document.add_paragraph(
        f"{yasmin_project.identity.land_title} {yasmin_project.identity.property_name} "
        "Pr?fecture"
    )
    document.save(damaged)

    with pytest.raises(GenerationError, match="encodage"):
        _verify(damaged, yasmin_project, "pv_division")


def test_pv_is_merged_in_the_requested_order_and_shared_fields_are_dynamic(tmp_path, yasmin_project):
    yasmin_project.identity.land_registry_office = "Conservation Atlas"
    yasmin_project.identity.prefecture = "Préfecture Atlas"
    yasmin_project.identity.commune = "Commune Oasis"
    DocumentGenerationService().generate(yasmin_project, tmp_path)

    pv_path = tmp_path / "PV_Division.docx"
    regulation_path = tmp_path / "Reglement_Copropriete.docx"
    pv = package_text(pv_path)
    regulation = package_text(regulation_path)
    assert pv.index("PROCES VERBAL DESCRIPTIF DE DIVISION") < pv.index("Après avoir reconnu")
    assert len(Document(pv_path).sections) >= 2
    assert "Conservation Atlas" in pv
    assert "Préfecture Atlas" in pv
    assert "Commune Oasis" in pv
    assert "L’an deux mille vingt-cinq, le lundi vingt octobre à dix heures." in pv
    assert "Karim El Mansouri" in pv and "AB123456" in pv
    assert "Karim El Mansouri" in regulation and "AB123456" in regulation
    assert "Préfecture de Préfecture Atlas" in regulation
    assert "Commune de Commune Oasis" in regulation
    assert yasmin_project.identity.overall_consistency in regulation
    for boundary in yasmin_project.identity.boundaries.values():
        assert boundary in regulation


def test_dynamic_level_blocks_remove_phantom_terrace_and_format_signed_cotes(tmp_path, yasmin_project):
    project = deepcopy(yasmin_project)
    project.levels = project.levels[:3]
    project.identity.overall_consistency = "Sous-sol + RDC + 1 étage"
    project.levels[0].start_elevation = Decimal("-2.20")
    project.levels[0].end_elevation = Decimal("0.40")
    DocumentGenerationService().generate(project, tmp_path)

    pv = Document(tmp_path / "PV_Division.docx")
    regulation = Document(tmp_path / "Reglement_Copropriete.docx")
    table_b = all_text(tmp_path / "Tableau_B.docx")
    assert "Terrasse" not in _heading_cells(pv)
    assert "Terrasse" not in _heading_cells(regulation)
    assert "PARTIE COMMUNE N° 7-7a" not in all_text(tmp_path / "PV_Division.docx")
    assert "PARTIE COMMUNE N° 7-7a" not in all_text(tmp_path / "Reglement_Copropriete.docx")
    for text in (all_text(tmp_path / "PV_Division.docx"), all_text(tmp_path / "Reglement_Copropriete.docx"), table_b):
        assert "-2,20" in text
        assert "+0,40" in text
        assert "+-" not in text


def test_multi_cotes_and_heights_are_written_consistently_in_documents(tmp_path, yasmin_project):
    project = deepcopy(yasmin_project)
    level = project.levels[0]
    level.end_elevations = (Decimal("3.10"), Decimal("5.70"))
    level.interior_heights = (Decimal("2.90"), Decimal("5.50"))

    DocumentGenerationService().generate(project, tmp_path)

    expected_range = "de la cote +0,20 m aux cotes +3,10 m et +5,70 m"
    expected_heights = "de hauteurs intérieures de 2,90 m et 5,50 m"
    pv_text = package_text(tmp_path / "PV_Division.docx")
    regulation_text = package_text(tmp_path / "Reglement_Copropriete.docx")
    table_b_text = package_text(tmp_path / "Tableau_B.docx")
    assert expected_range in pv_text
    assert expected_range in regulation_text
    assert expected_range in table_b_text
    assert expected_heights in pv_text
    assert expected_heights in regulation_text
    for filename in EXPECTED_OUTPUTS:
        assert_french_characters_are_intact(tmp_path / filename)


def test_tableau_b_private_suffixes_match_tableau_a_and_common_rows_are_unsuffixed(tmp_path, yasmin_project):
    DocumentGenerationService().generate(yasmin_project, tmp_path)
    table_a = Document(tmp_path / "Tableau_A.docx").tables[0]
    expected_names = [row.cells[0].text.strip() for row in table_a.rows[1:]]
    table_b = Document(tmp_path / "Tableau_B.docx").tables[0]
    private_names = []
    common_names = []
    private_indices = {
        part.index for level in yasmin_project.levels for part in level.parts
        if part.nature.value == "privative"
    }
    common_indices = {
        part.index for level in yasmin_project.levels for part in level.parts
        if part.nature.value == "commune"
    }
    for row in table_b.rows:
        cells = row.cells
        if len(cells) < 4:
            continue
        if cells[2].text.strip() in private_indices:
            private_names.append(cells[0].text.strip())
        elif cells[3].text.strip() in common_indices:
            common_names.append(cells[0].text.strip())
    assert private_names == expected_names
    assert common_names
    assert set(common_names) == {yasmin_project.identity.property_name}


def test_prefecture_is_written_under_table_logos_and_in_recap(tmp_path, yasmin_project):
    yasmin_project.identity.prefecture = "Préfecture Atlas"
    yasmin_project.identity.commune = "Commune Oasis"
    DocumentGenerationService().generate(yasmin_project, tmp_path)
    for filename in ("Tableau_A.docx", "Tableau_B.docx", "Tableau_Recapitulatif.docx"):
        text = package_text(tmp_path / filename)
        assert "Préfecture Atlas" in text
        assert "Meknès Al Ismaïlia" not in text


def test_legal_text_and_calculation_tables_remain_present(tmp_path, yasmin_project):
    DocumentGenerationService().generate(yasmin_project, tmp_path)
    regulation = Document(tmp_path / "Reglement_Copropriete.docx")
    recap = Document(tmp_path / "Tableau_Recapitulatif.docx")
    table_b = Document(tmp_path / "Tableau_B.docx")
    assert len(table_b.tables[0].rows) == 25
    distribution = next(
        table for table in regulation.tables
        if any("Indice des parties" in cell.text for row in table.rows for cell in row.cells)
    )
    assert len(distribution.rows) == 27
    assert [len(table.rows) for table in recap.tables] == [1, 6, 1, 4]
    regulation_rows = [[cell.text for cell in row.cells] for row in distribution.rows]
    assert any("225" in row and "247" in row and "10000" in row for row in regulation_rows)
    reference = Document(resource_root() / "templates" / "runtime" / "reglement.docx")
    is_article_one = lambda text: text.replace("\xa0", " ").strip().startswith("Article 1 :")
    reference_article = next(p.text for p in reference.paragraphs if is_article_one(p.text))
    generated_article = next(p.text for p in regulation.paragraphs if is_article_one(p.text))
    assert generated_article == reference_article
    assert regulation.paragraphs[-1].text != ""
    footer_paragraph = regulation.sections[0].footer.paragraphs[0]
    assert footer_paragraph.paragraph_format.right_indent.inches == pytest.approx(0.75)


def test_generated_documents_request_field_refresh(tmp_path, yasmin_project):
    DocumentGenerationService().generate(yasmin_project, tmp_path)
    namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    for filename in EXPECTED_OUTPUTS:
        with ZipFile(tmp_path / filename) as archive:
            settings = etree.fromstring(archive.read("word/settings.xml"))
        values = settings.xpath("./w:updateFields/@w:val", namespaces=namespaces)
        assert values == ["true"]


def test_merged_pv_preserves_every_referenced_style_and_numbering_definition():
    template = resource_root() / "templates" / "runtime" / "pv_division.docx"
    assert len(Document(template).sections) == 2
    namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(template) as archive:
        styles = etree.fromstring(archive.read("word/styles.xml"))
        numbering = etree.fromstring(archive.read("word/numbering.xml"))
        defined_styles = set(styles.xpath(".//w:style/@w:styleId", namespaces=namespaces))
        defined_numbers = set(numbering.xpath(".//w:num/@w:numId", namespaces=namespaces))
        used_styles = set()
        used_numbers = set()
        for name in archive.namelist():
            if not (name.startswith("word/") and name.endswith(".xml")):
                continue
            root = etree.fromstring(archive.read(name))
            used_styles.update(root.xpath(
                ".//w:pStyle/@w:val | .//w:rStyle/@w:val", namespaces=namespaces,
            ))
            used_numbers.update(root.xpath(".//w:numId/@w:val", namespaces=namespaces))
    assert used_styles <= defined_styles
    assert used_numbers <= defined_numbers


def test_level_description_does_not_duplicate_existing_article(yasmin_project):
    level = yasmin_project.levels[1]
    level.name = "Le Premier Étage"
    description = level_description(level)
    assert description.startswith("Le Premier Étage couvrant")
    assert "Le Le" not in description


def test_dynamic_part_runs_retain_the_template_font_size(tmp_path, yasmin_project):
    DocumentGenerationService().generate(yasmin_project, tmp_path)
    regulation = Document(tmp_path / "Reglement_Copropriete.docx")
    part_paragraph = next(
        paragraph
        for table in regulation.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
        if paragraph.text.startswith("PARTIE PRIVATIVE N° 1")
    )
    assert part_paragraph.runs[0].font.size.pt == 12
    assert part_paragraph.runs[1].font.size.pt == 12
